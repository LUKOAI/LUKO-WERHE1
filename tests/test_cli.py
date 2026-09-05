import pytest
from openpyxl import load_workbook

from amazon_vat_merger.cli import main
from tests.conftest import SAMPLE_CSV, SAMPLE_PDF_DIR


@pytest.mark.skipif(not SAMPLE_CSV.exists(), reason="brak próbki CSV")
def test_cli_end_to_end(tmp_path, capsys):
    out = tmp_path / "wynik.xlsx"
    js = tmp_path / "faktury.json"
    rc = main(["--csv", str(SAMPLE_CSV), "--pdf", str(SAMPLE_PDF_DIR), "--out", str(out), "--json", str(js), "--no-nbp"])
    assert rc == 0 and out.exists() and js.exists()
    wb = load_workbook(out)
    assert wb.sheetnames[0] == "Wszystko" and wb.sheetnames[-1] == "Diagnostyka"
    ws = wb["Wszystko"]
    header = [c.value for c in ws[1]]
    assert header[:3] == ["Zakładka", "Kategoria", "Numer faktury VAT"]
    assert ws.max_row == 45  # 44 transakcje + nagłówek
    # faktura CZ dopasowana, z nazwiskiem z PDF
    col_inv = header.index("Numer faktury VAT") + 1
    col_name = header.index("Imię i nazwisko Kupującego") + 1
    col_match = header.index("Dopasowanie PDF") + 1
    rows = {ws.cell(r, col_inv).value: r for r in range(2, ws.max_row + 1)}
    r = rows["CZ60009FG6O6HI"]
    assert ws.cell(r, col_name).value == "Stylem sp. z o.o."
    assert ws.cell(r, col_match).value == "PDF"
    n_pdf = len(list(SAMPLE_PDF_DIR.glob("*.pdf")))
    matched = sum(1 for r in range(2, ws.max_row + 1) if ws.cell(r, col_match).value == "PDF")
    assert matched == min(n_pdf, 44)
    # zakładki grup
    assert "DE OSS" in wb.sheetnames and "CZ WDT" in wb.sheetnames and "FR Marketplace" in wb.sheetnames
    de = wb["DE OSS"]
    assert de["A1"].value is None and de["B1"].value == "Niemcy" and de["C1"].value == "OSS"
    assert all(c.value is None for c in de[3])           # pusty wiersz 3 jak w arkuszu klienta
    last = de.max_row
    assert de.cell(last, 1).value == "RAZEM"
    hdr = [c.value for c in de[2]]
    ci = hdr.index("Kwota netto EUR") + 1
    assert str(de.cell(last, ci).value).startswith("=SUM(")
    assert de.cell(4, hdr.index("Stawka VAT") + 1).number_format == "0%"
    printed = capsys.readouterr().out
    assert f"dopasowane: {matched}" in printed
