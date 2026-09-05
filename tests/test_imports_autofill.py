from datetime import date
from pathlib import Path

from fbaplan.catalog.store import Resolver, load_products
from fbaplan.history.store import compute_stats
from fbaplan.imports import export_catalog_xlsx, import_catalog_xlsx, import_fee_preview, import_inventory_report, read_table
from fbaplan.models import Alias, CartonSpec, Dims, HistoryLine, Plan, PlanLine, Product, StockRow
from fbaplan.planner.autofill import autofill
from fbaplan.predict.fc import FCPredictor
from fbaplan.stock import load_stock


def _products():
    return {
        "1ZWG80dp": Product("1ZWG80dp", "Świder do Gleby 80D SDS PLUS", "auger"),
        "4APlus": Product("4APlus", "Adapter do świdra SDS PLUS", "adapter", mskus={"WERHE": "WH-ADAPTER-PLUS"}),
    }


def _aliases():
    return [Alias("WERHE Wiertło do ziemi 80 mm z SDS Plus rejestrujące precyzyjne wiertło", "1ZWG80dp", "WERHE", 0.9, "t")]


def test_fee_preview_import(tmp_path: Path):
    f = tmp_path / "fee.txt"
    f.write_text(
        "sku\tfnsku\tasin\tproduct-name\tlongest-side\tmedian-side\tshortest-side\tunit-of-dimension\titem-package-weight\tunit-of-weight\tproduct-size-tier\n"
        "WH-80-SDSP\tX1\tB01\tWERHE Wiertło do ziemi 80 mm z SDS Plus rejestrujące precyzyjne wiertło\t82.5\t10.2\t9.8\tcentimeters\t2.9\tkilograms\tStandard-Übergröße\n"
        "WH-ADAPTER-PLUS\tX2\tB02\tWERHE Adapter SDS Plus\t4.7\t2.0\t1.9\tinches\t0.45\tpounds\tKleines Paket\n"
        "WH-UNKNOWN\tX3\tB03\tCoś zupełnie innego 12 cm\t10\t5\t5\tcentimeters\t0.1\tkilograms\tKoperta\n",
        encoding="utf-8",
    )
    products = _products()
    rep = import_fee_preview(f, products, Resolver(products, _aliases()), tmp_path / "products.csv")
    assert rep.rows == 3 and rep.matched == 2 and len(rep.unmatched) == 1
    p = products["1ZWG80dp"]
    assert p.unit_dims == Dims(82.5, 10.2, 9.8) and p.unit_weight_kg == 2.9
    assert p.mskus["WERHE"] == "WH-80-SDSP"
    a = products["4APlus"]
    assert a.unit_dims is not None and abs(a.unit_dims.length_cm - 11.94) < 0.01 and abs(a.unit_weight_kg - 0.204) < 0.001
    assert "Kleines Paket" in a.notes
    saved = load_products(tmp_path / "products.csv")
    assert saved["1ZWG80dp"].unit_dims == Dims(82.5, 10.2, 9.8)


def test_inventory_and_restock_import_merge(tmp_path: Path):
    inv = tmp_path / "inventory.txt"
    inv.write_text(
        "sku\tfnsku\tasin\tproduct-name\tafn-fulfillable-quantity\tafn-inbound-shipped-quantity\tafn-inbound-working-quantity\tafn-inbound-receiving-quantity\n"
        "WH-80-SDSP\tX1\tB01\tWERHE Wiertło do ziemi 80 mm z SDS Plus rejestrujące precyzyjne wiertło\t12\t20\t0\t5\n"
        "WH-ADAPTER-PLUS\tX2\tB02\tWERHE Adapter SDS Plus\t300\t0\t0\t0\n",
        encoding="cp1250",
    )
    restock = tmp_path / "restock.csv"
    restock.write_text(
        "Merchant SKU;Product Name;Units Sold Last 30 Days;Available;Inbound\n"
        "WH-80-SDSP;WERHE Wiertło do ziemi 80 mm z SDS Plus rejestrujące precyzyjne wiertło;90;12;25\n",
        encoding="utf-8-sig",
    )
    products = _products()
    r = Resolver(products, _aliases())
    stock_path = tmp_path / "stock.csv"
    rep1 = import_inventory_report(inv, products, r, stock_path)
    assert rep1.matched == 2
    rep2 = import_inventory_report(restock, products, r, stock_path)
    assert rep2.matched == 1
    stock = load_stock(stock_path, products)
    assert stock["1ZWG80dp"].fba_available == 12 and stock["1ZWG80dp"].fba_inbound == 25 and stock["1ZWG80dp"].sales_30d == 90
    assert stock["4APlus"].fba_available == 300
    assert stock["1ZWG80dp"].cover_days is not None and 11 < stock["1ZWG80dp"].cover_days < 13


def test_catalog_xlsx_roundtrip(tmp_path: Path):
    products = _products()
    x = tmp_path / "katalog.xlsx"
    n = export_catalog_xlsx(x, products.values(), {"1ZWG80dp": 500}, {"1ZWG80dp": "XPO1", "4APlus": "WRO5"})
    assert n == 2
    rows = read_table(x)
    assert rows[0]["sku"] == "1ZWG80dp"  # sorted by priority
    # simulate the client filling the workbook
    from openpyxl import load_workbook

    wb = load_workbook(x)
    ws = wb["Katalog"]
    hdr = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(hdr)}
    ws.cell(row=2, column=col["Sztuka dł. [cm]"], value=82)
    ws.cell(row=2, column=col["Sztuka szer. [cm]"], value=10)
    ws.cell(row=2, column=col["Sztuka wys. [cm]"], value=10)
    ws.cell(row=2, column=col["Sztuka waga [kg]"], value="2,9")
    ws.cell(row=2, column=col["Szt. w kartonie"], value=11)
    ws.cell(row=2, column=col["Karton dł. [cm]"], value=84)
    ws.cell(row=2, column=col["Karton szer. [cm]"], value=33)
    ws.cell(row=2, column=col["Karton wys. [cm]"], value=22)
    ws.cell(row=2, column=col["Grupa pakowania"], value="dlugie")
    ws.cell(row=3, column=col["Aktywny (1/0)"], value=0)
    wb.save(x)
    rep = import_catalog_xlsx(x, products, tmp_path / "products.csv")
    assert rep.updated == 2
    p = products["1ZWG80dp"]
    assert p.unit_dims == Dims(82, 10, 10) and p.unit_weight_kg == 2.9
    assert p.carton.units_per_carton == 11 and p.carton.dims == Dims(84, 33, 22) and not p.carton.estimated
    assert p.pack_group == "dlugie" and products["4APlus"].active is False


def test_autofill_fills_pallet_same_fc():
    cat = {
        "A": Product("A", "Świder 80x800", "auger", unit_dims=Dims(82, 10, 10), unit_weight_kg=3.0, carton=CartonSpec(4, Dims(84, 22, 22), 12.5)),
        "B": Product("B", "Świder 100x800", "auger", unit_dims=Dims(82, 12, 12), unit_weight_kg=3.6, carton=CartonSpec(4, Dims(84, 26, 26), 15.0)),
        "C": Product("C", "Dłuto 75x600 SDS Max", "chisel_flat", unit_dims=Dims(61, 8, 4), unit_weight_kg=2.0, carton=CartonSpec(10, Dims(63, 30, 20), 20.5)),
        "D": Product("D", "Adapter SDS Plus", "adapter", unit_dims=Dims(12, 5, 5), unit_weight_kg=0.2, carton=CartonSpec(50, Dims(40, 30, 20), 10.5)),
    }
    hist = []
    for _ in range(5):
        hist += [HistoryLine("A", "XPO1", date(2026, 5, 1), 20, "S1"), HistoryLine("B", "XPO1", date(2026, 5, 1), 12, "S1"),
                 HistoryLine("C", "XPO1", date(2026, 5, 1), 10, "S2"), HistoryLine("D", "WRO5", date(2026, 5, 1), 50, "S3")]
    pr = FCPredictor(hist, as_of=date(2026, 9, 1))
    stats = compute_stats(hist, as_of=date(2026, 9, 1))
    stock = {"B": StockRow("B", fba_available=2, sales_30d=50), "C": StockRow("C", fba_available=10, sales_30d=40), "D": StockRow("D", fba_available=0, sales_30d=500)}
    plan = Plan("P", "2026-09-01T10:00:00", mode="pallet", max_pallets=1)
    plan.lines = [PlanLine("A", 8)]
    res = autofill(plan, cat, pr, stock, stats, target_fill=0.6)
    assert res.added, res.stopped_because
    skus = {s for s, _ in res.added}
    assert "D" not in skus and skus <= {"A", "B", "C"}
    ev = res.final
    assert not ev.split and len(ev.groups["XPO1"].pallets) == 1
    assert ev.groups["XPO1"].pallets[0].fill_ratio >= 0.6
    for l in plan.lines:
        assert l.qty % cat[l.sku].units_per_carton == 0


def test_autofill_without_dims_uses_units_target():
    cat = {"A": Product("A", "Świder 80x800", "auger", carton=CartonSpec(11)), "B": Product("B", "Świder 100x800", "auger", carton=CartonSpec(10))}
    hist = [HistoryLine("A", "XPO1", date(2026, 5, 1), 22, "S1"), HistoryLine("B", "XPO1", date(2026, 5, 1), 10, "S1")] * 4
    pr = FCPredictor(hist, as_of=date(2026, 9, 1))
    plan = Plan("P", "2026-09-01T10:00:00", mode="pallet")
    plan.lines = [PlanLine("A", 22)]
    res = autofill(plan, cat, pr, None, compute_stats(hist, as_of=date(2026, 9, 1)), target_units=60)
    assert res.final.units >= 60 and "szt." in res.stopped_because
