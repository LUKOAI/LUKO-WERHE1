"""Zapis arkuszy do pliku .xlsx (openpyxl)."""
from __future__ import annotations

import os
import re
from datetime import date
from pathlib import Path

from .merge import Formula, Link, Sheet

_BAD_SHEET_CHARS = re.compile(r"[\[\]:*?/\\]")


def safe_sheet_name(name: str, used: set[str]) -> str:
    base = _BAD_SHEET_CHARS.sub("_", name).strip() or "Arkusz"
    base = base[:31]
    candidate = base
    n = 2
    while candidate.lower() in used:
        suffix = f" ({n})"
        candidate = base[: 31 - len(suffix)] + suffix
        n += 1
    used.add(candidate.lower())
    return candidate


def write_xlsx(path: str | Path, sheets: list[Sheet]) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    used: set[str] = set()
    # nazwy zakładek najpierw dla wszystkich – linki z „Wszystko” muszą wskazywać nazwę faktycznie użytą w pliku
    titles = {sheet.name: safe_sheet_name(sheet.name, used) for sheet in sheets}
    bold = Font(bold=True)
    link_font = Font(color="0563C1", underline="single")
    fill = PatternFill("solid", fgColor="DDEBF7")
    for sheet in sheets:
        ws = wb.create_sheet(titles[sheet.name])
        for ri, row in enumerate(sheet.rows, start=1):
            for ci, val in enumerate(row, start=1):
                if val is None:
                    continue
                cell = ws.cell(row=ri, column=ci)
                if isinstance(val, Link):
                    cell.value = Link.excel_formula(titles.get(val.tab, val.tab), val.row, val.text, val.key)
                    cell.font = link_font
                elif isinstance(val, Formula):
                    cell.value = str(val)
                elif isinstance(val, date):
                    cell.value = val
                    cell.number_format = "yyyy-mm-dd"
                else:
                    cell.value = val
                    if isinstance(val, str) and val.startswith("="):
                        cell.data_type = "s"   # tekst zaczynający się od '=' to nie formuła
                if ri > sheet.header_row:
                    if ci in sheet.money_cols and isinstance(val, (int, float, Formula)):
                        cell.number_format = "#,##0.00"
                    elif ci in sheet.pct_cols and isinstance(val, (int, float)):
                        cell.number_format = "0.0%"   # stawki obniżone (5,5 %, 13,5 %) nie mogą się zaokrąglać
        # nagłówek
        fill_extra = PatternFill("solid", fgColor="EDEDED")
        for ci in range(1, len(sheet.rows[sheet.header_row - 1]) + 1 if sheet.rows else 1):
            c = ws.cell(row=sheet.header_row, column=ci)
            c.font = bold
            c.fill = fill_extra if (sheet.extras_from and ci >= sheet.extras_from) else fill
            c.alignment = Alignment(wrap_text=True, vertical="top")
        if sheet.header_row > 1:
            for ci in range(1, len(sheet.rows[0]) + 1):
                ws.cell(row=1, column=ci).font = Font(bold=True, size=12)
        last = ws.max_row
        if sheet.rows and isinstance(sheet.rows[-1][0], str) and sheet.rows[-1][0] == "RAZEM":
            for ci in range(1, len(sheet.rows[-1]) + 1):
                ws.cell(row=last, column=ci).font = bold
        ws.freeze_panes = ws.cell(row=sheet.header_row + 1, column=2)
        # szerokości kolumn
        ncols = max((len(r) for r in sheet.rows), default=0)
        for ci in range(1, ncols + 1):
            width = 10
            for row in sheet.rows[sheet.header_row - 1: sheet.header_row + 40]:
                v = row[ci - 1] if ci - 1 < len(row) else None
                if v is None:
                    continue
                # link mierzymy po widocznej etykiecie, formułę stałą szerokością – nie po tekście formuły
                shown = v.text if isinstance(v, Link) else ("0000000.00" if isinstance(v, Formula) else str(v))
                width = max(width, min(60, len(shown) + 2))
            ws.column_dimensions[get_column_letter(ci)].width = width
        filter_last = last - 1 if (sheet.rows and sheet.rows[-1][0] == "RAZEM") else last
        ws.auto_filter.ref = f"A{sheet.header_row}:{get_column_letter(max(ncols, 1))}{max(filter_last, sheet.header_row)}"
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    # zapis do pliku tymczasowego i podmiana – przerwany zapis nie zostawia uszkodzonego pliku pod docelową nazwą
    tmp = path.with_name(path.name + ".tmp")
    try:
        wb.save(str(tmp))
        os.replace(tmp, path)
    except PermissionError as exc:
        raise PermissionError(
            f"nie można zapisać {path} – plik jest otwarty w innym programie (Excel)? "
            f"Zamknij go i uruchom ponownie. ({exc})") from exc
    finally:
        if tmp.exists():
            tmp.unlink(missing_ok=True)
    return path
