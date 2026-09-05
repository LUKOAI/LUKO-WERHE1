"""Parsery miesiecznych plikow z planami wysylki FBA.

Formaty spotykane w archiwum klienta (2020-2026):

* ``.docx`` (2020-09 .. 2022-01): akapity tekstu::

      FBA (01.12.21, 09:22) - 1
      FBA15FK58H6D
      <nazwa produktu> – <ilosc>
      <nazwa produktu> – <ilosc>

* ``.xlsx`` (2022-02 .. 2026-06): arkusz z kolumnami
  ``Numer wysyłki | produkt | Ilość``; pierwszy wiersz wysylki ma naglowek
  ``FBA STA (01/06/2026 06:42)-WRO5`` w kol. A, kolejny wiersz ma ID
  ``FBA15LV4JK05 | 4QZVZJWE`` w kol. A; produkty w kol. B, ilosci w kol. C.

* ``.pdf`` (2026-07 ..): eksport/wydruk tabeli - obslugiwany przez pdfplumber.

Wynik: lista :class:`Shipment` z liniami :class:`Line`.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Iterable, Iterator, Optional

# -- wzorce -----------------------------------------------------------------

# FBA (01.12.21, 09:22) - 1        (docx)
# FBA STA (01/06/2026 06:42)-WRO5  (xlsx)
HEADER_RE = re.compile(
    r"^\s*FBA\s*(?P<kind>STA)?\s*\(\s*"
    r"(?P<d>\d{1,2})[./](?P<m>\d{1,2})[./](?P<y>\d{2,4})"
    r"\s*,?\s*(?P<hh>\d{1,2}):(?P<mm>\d{2})\s*\)"
    r"\s*[-–]?\s*(?P<suffix>[A-Za-z0-9]+)?\s*$"
)
# FBA15FK58H6D  |  FBA15LV4JK05 | 4QZVZJWE
ID_RE = re.compile(r"^\s*(?P<id>FBA[0-9A-Z]{9,})\s*(?:[|,;]\s*(?P<ref>[0-9A-Z]{6,}))?\s*$")
# "<nazwa> – 12"  (ostatni myslnik przed liczba na koncu)
LINE_QTY_RE = re.compile(r"^(?P<name>.*?)\s*[–—-]\s*(?P<qty>\d+)\s*$")
LINE_NOQTY_RE = re.compile(r"^(?P<name>.+?)\s*[–—-]\s*$")
QTY_ONLY_RE = re.compile(r"^\s*[–—-]?\s*(?P<qty>\d{1,4})\s*$")
FC_RE = re.compile(r"^[A-Z]{3,4}\d{0,2}$")


@dataclass
class Line:
    name: str
    qty: Optional[int]
    raw: str = ""
    skus: list[str] = field(default_factory=list)  # wewnetrzne kody SKU (wydruk Apilo)


@dataclass
class Shipment:
    source_file: str
    created_at: Optional[datetime]
    shipment_id: Optional[str] = None
    reference_id: Optional[str] = None
    fc: Optional[str] = None
    header_suffix: Optional[str] = None
    lines: list[Line] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)  # wolny tekst z pliku (np. uwagi pracownika)

    @property
    def units(self) -> int:
        return sum(l.qty or 0 for l in self.lines)


def _parse_header(text: str) -> Optional[dict]:
    m = HEADER_RE.match(text)
    if not m:
        return None
    y = int(m.group("y"))
    if y < 100:
        y += 2000
    try:
        dt = datetime(y, int(m.group("m")), int(m.group("d")), int(m.group("hh")), int(m.group("mm")))
    except ValueError:
        dt = None
    suffix = m.group("suffix")
    fc = suffix if suffix and FC_RE.match(suffix) else None
    return {"created_at": dt, "suffix": suffix, "fc": fc, "kind": m.group("kind")}


def _parse_line(text: str) -> Optional[Line]:
    t = text.strip()
    if not t:
        return None
    m = LINE_QTY_RE.match(t)
    if m:
        return Line(name=m.group("name").strip(), qty=int(m.group("qty")), raw=t)
    m = LINE_NOQTY_RE.match(t)
    if m:
        return Line(name=m.group("name").strip(), qty=None, raw=t)
    return Line(name=t, qty=None, raw=t)


# -- docx -------------------------------------------------------------------

def _docx_lines(path: Path) -> Iterator[str]:
    import docx  # python-docx

    d = docx.Document(str(path))
    for p in d.paragraphs:
        for part in p.text.split("\n"):
            yield part
    for t in d.tables:
        for r in t.rows:
            for c in r.cells:
                for part in c.text.split("\n"):
                    yield part


def _split_header_and_id(text: str) -> list[str]:
    """'FBA (01.09.20, 12:30) – 1\t\tFBA15D52DVJR' -> ['FBA (01.09.20, 12:30) – 1', 'FBA15D52DVJR']"""
    parts = [t.strip() for t in re.split(r"\t+|\s{3,}", text) if t.strip()]
    if len(parts) >= 2 and _parse_header(parts[0]) and ID_RE.match(parts[-1]):
        return [parts[0], parts[-1]]
    if len(parts) >= 2 and ID_RE.match(parts[0]) and _parse_header(parts[-1]):
        # reversed order: 'FBA15DNJB6W5\t\tFBA (22.03.21, 07:59) – 1'
        return [parts[-1], parts[0]]
    m = re.match(r"^(?P<h>FBA\s*\(.*?\)\s*[-–]?\s*\w*)\s+(?P<id>FBA[0-9A-Z]{9,})\s*$", text)
    if m:
        return [m.group("h"), m.group("id")]
    return [text]


def parse_docx(path: Path) -> list[Shipment]:
    shipments: list[Shipment] = []
    cur: Optional[Shipment] = None
    pending: list[str] = []
    for raw in _docx_lines(path):
        pending.extend(_split_header_and_id(raw.strip()) if raw.strip() else [])
    for text in pending:
        text = text.strip()
        if not text:
            continue
        qm = QTY_ONLY_RE.match(text)
        if qm and cur is not None and cur.lines and cur.lines[-1].qty is None:
            cur.lines[-1].qty = int(qm.group("qty"))
            cur.warnings = [w for w in cur.warnings if not w.startswith("qty-missing: " + cur.lines[-1].raw[:60])]
            continue
        h = _parse_header(text)
        if h:
            cur = Shipment(source_file=path.name, created_at=h["created_at"], fc=h["fc"], header_suffix=h["suffix"])
            shipments.append(cur)
            continue
        idm = ID_RE.match(text)
        if idm:
            if cur is None:
                cur = Shipment(source_file=path.name, created_at=None)
                cur.warnings.append("id-before-header")
                shipments.append(cur)
            if cur.shipment_id:
                # drugi ID pod tym samym naglowkiem -> nowa wysylka bez daty
                nxt = Shipment(source_file=path.name, created_at=cur.created_at, fc=cur.fc, header_suffix=cur.header_suffix)
                nxt.warnings.append("id-without-own-header")
                shipments.append(nxt)
                cur = nxt
            cur.shipment_id = idm.group("id")
            cur.reference_id = idm.group("ref")
            continue
        line = _parse_line(text)
        if line is None:
            continue
        if cur is None:
            cur = Shipment(source_file=path.name, created_at=None)
            cur.warnings.append("lines-before-header")
            shipments.append(cur)
        if line.qty is None:
            cur.warnings.append(f"qty-missing: {line.raw[:60]}")
        cur.lines.append(line)
    return shipments


# -- xlsx -------------------------------------------------------------------

def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _to_int(v) -> Optional[int]:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return int(v)
    s = str(v).strip().replace(",", ".")
    m = re.search(r"-?\d+(?:\.\d+)?", s)
    return int(float(m.group())) if m else None


def parse_xlsx(path: Path) -> list[Shipment]:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    shipments: list[Shipment] = []
    for ws in wb.worksheets:
        cur: Optional[Shipment] = None
        for row in ws.iter_rows(values_only=True):
            cells = [_cell(c) for c in row] + ["", "", ""]
            a, b, c = cells[0], cells[1], row[2] if len(row) > 2 else None
            if a:
                h = _parse_header(a)
                if h:
                    cur = Shipment(source_file=path.name, created_at=h["created_at"], fc=h["fc"], header_suffix=h["suffix"])
                    shipments.append(cur)
                else:
                    idm = ID_RE.match(a)
                    if idm and cur is not None:
                        if cur.shipment_id and cur.shipment_id != idm.group("id"):
                            nxt = Shipment(source_file=path.name, created_at=cur.created_at, fc=cur.fc, header_suffix=cur.header_suffix)
                            nxt.warnings.append("id-without-own-header")
                            shipments.append(nxt)
                            cur = nxt
                        cur.shipment_id = idm.group("id")
                        cur.reference_id = idm.group("ref")
                    elif idm:
                        cur = Shipment(source_file=path.name, created_at=None, shipment_id=idm.group("id"), reference_id=idm.group("ref"))
                        cur.warnings.append("id-before-header")
                        shipments.append(cur)
                    elif a.lower().startswith("numer wysy"):
                        continue  # naglowek tabeli
                    else:
                        if cur is not None:
                            cur.notes.append(a)
            if b:
                if cur is None:
                    cur = Shipment(source_file=path.name, created_at=None)
                    cur.warnings.append("lines-before-header")
                    shipments.append(cur)
                if b.lower() in ("produkt", "product"):
                    continue
                qty = _to_int(c)
                # czasem ilosc doklejona do nazwy "nazwa – 12"
                if qty is None:
                    ln = _parse_line(b)
                    if ln and ln.qty is not None:
                        cur.lines.append(ln)
                        continue
                    cur.warnings.append(f"qty-missing: {b[:60]}")
                cur.lines.append(Line(name=b, qty=qty, raw=f"{b} | {c}"))
    return shipments


# -- pdf (wydruk z Apilo) -----------------------------------------------------

# "1. PF260806243 | Palety Amazon"  /  "2. AF260806233 | Paczki FBA"
APILO_HDR_RE = re.compile(r"^\s*\d+\.\s+(?P<ref>[A-Z]{1,3}\d{6,})\s*\|\s*(?P<type>.+?)\s*$")
# "Dane wysyłki: WRO5 - Okmiany Chojnow - Amazon | Amazonska WRO5 | 59225 CHOJNÓW | Polska | ..."
APILO_DEST_RE = re.compile(r"^\s*Dane wysy\S+:\s*(?P<fc>[A-Z]{2,4}\d)\s*-\s*(?P<rest>.*)$")
# "3. Wbijak do pali 85x105x234mm SDS Max / 6PO85x234m 48"
APILO_ITEM_RE = re.compile(r"^\s*(?P<lp>\d+)\.\s+(?P<name>.+?)\s+(?P<qty>\d+)\s*$")
APILO_SKU_SPLIT_RE = re.compile(r"\s+/\s+")
# kod SKU klienta: cyfra rodziny + litery, np. 1ZWG60dp, 3DB75x410h30, 4AM14, 6PO13,5m, 7NT744D
APILO_SKU_RE = re.compile(r"^\d[A-Za-z][\w.,/-]*$")


def _apilo_date(ref: str) -> Optional[datetime]:
    """Numer Apilo = prefiks + RRMM + licznik (np. PF260806243 -> 2026-08); dzien nieznany -> 1."""
    m = re.match(r"^[A-Z]+(\d{2})(\d{2})\d+$", ref)
    if not m:
        return None
    try:
        return datetime(2000 + int(m.group(1)), int(m.group(2)), 1)
    except ValueError:
        return None


def split_apilo_name(name: str) -> tuple[str, list[str]]:
    """'Dłuto 75x410 / 3DB75x410h30 / 6PO85x220h30' -> ('Dłuto 75x410', ['3DB75x410h30', ...])."""
    parts = APILO_SKU_SPLIT_RE.split(name.strip())
    if len(parts) == 1:
        return parts[0], []
    base, skus = parts[0], []
    for p in parts[1:]:
        if APILO_SKU_RE.match(p):
            skus.append(p)
        else:
            base = f"{base} / {p}"
    return base.strip(), skus


def parse_pdf(path: Path) -> list[Shipment]:
    import pdfplumber

    shipments: list[Shipment] = []
    cur: Optional[Shipment] = None
    with pdfplumber.open(str(path)) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            tables = [t for t in (page.extract_tables() or []) if t]
            # naglowki + tabele w kolejnosci wystapienia na stronie
            headers: list[tuple[str, str, Optional[str]]] = []  # (ref, type, fc)
            pending_fc: Optional[str] = None
            for ln in text.split("\n"):
                hm = APILO_HDR_RE.match(ln)
                if hm:
                    headers.append([hm.group("ref"), hm.group("type"), None])
                    continue
                dm = APILO_DEST_RE.match(ln)
                if dm and headers:
                    headers[-1][2] = dm.group("fc")
            # tabela kontynuowana z poprzedniej strony (brak naglowka przed pierwsza tabela)
            first_line = next((l for l in text.split("\n") if l.strip()), "")
            continued = bool(tables) and not APILO_HDR_RE.match(first_line) and cur is not None
            ti = 0
            if continued and len(tables) > len(headers):
                _append_apilo_rows(cur, tables[0])
                ti = 1
            for h in headers:
                ref, typ, fc = h
                cur = Shipment(source_file=path.name, created_at=_apilo_date(ref), shipment_id=None,
                               reference_id=ref, fc=fc, header_suffix=typ)
                shipments.append(cur)
                if ti < len(tables):
                    _append_apilo_rows(cur, tables[ti])
                    ti += 1
                else:
                    cur.warnings.append("no-table-for-header")
            if ti < len(tables):
                # nadmiarowe tabele - doklej do ostatniej wysylki
                for t in tables[ti:]:
                    if cur is not None:
                        _append_apilo_rows(cur, t)
                        cur.warnings.append("extra-table-appended")
    return shipments


def _append_apilo_rows(sh: Shipment, table) -> None:
    for r in table:
        cells = [_cell(x) for x in (r or [])] + ["", "", ""]
        lp, name, qty = cells[0], cells[1], cells[2]
        if not name or name.lower() == "nazwa":
            continue
        base, skus = split_apilo_name(name)
        q = _to_int(qty)
        if q is None:
            sh.warnings.append(f"qty-missing: {name[:60]}")
        ln = Line(name=base, qty=q, raw=name, skus=skus)
        sh.lines.append(ln)


# -- API --------------------------------------------------------------------

PARSERS = {".docx": parse_docx, ".xlsx": parse_xlsx, ".pdf": parse_pdf}


def parse_file(path: Path | str) -> list[Shipment]:
    path = Path(path)
    fn = PARSERS.get(path.suffix.lower())
    if fn is None:
        raise ValueError(f"Nieobslugiwany format: {path.suffix}")
    return fn(path)


def parse_many(paths: Iterable[Path | str]) -> list[Shipment]:
    out: list[Shipment] = []
    for p in paths:
        out.extend(parse_file(p))
    return out
