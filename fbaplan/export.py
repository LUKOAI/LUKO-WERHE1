"""Exports: Send-to-Amazon line list, packing list per carton / pallet, summary."""
from __future__ import annotations

import csv
import io
from pathlib import Path
from typing import Optional

from .planner.session import PlanEval


def send_to_amazon_rows(ev: PlanEval, fc: Optional[str] = None) -> list[dict]:
    """Rows (merchant SKU, quantity) for the Seller Central 'Send to Amazon' step, optionally for one FC."""
    rows = []
    for le in ev.lines:
        if fc and le.prediction.fc != fc:
            continue
        p = le.product
        msku = le.line.msku or (p.msku_for(le.line.brand) if p else le.line.sku)
        rows.append({"MerchantSKU": msku, "Quantity": le.line.qty, "SKU": le.line.sku,
                     "Nazwa": p.name if p else "", "Magazyn (przewidywany)": le.prediction.fc or "?"})
    return rows


def packing_rows(ev: PlanEval) -> list[dict]:
    rows = []
    for fc, g in ev.groups.items():
        for pi, pal in enumerate(g.pallets, 1):
            for li, layer in enumerate(pal.layers, 1):
                for ci, c in enumerate(layer.cartons, 1):
                    for sku, q in c.contents.items():
                        p = ev.lines and next((le.product for le in ev.lines if le.line.sku == sku), None)
                        rows.append({
                            "Magazyn": fc, "Paleta": pi, "Warstwa": li, "Karton": ci, "SKU": sku,
                            "Nazwa": p.name if p else "", "Sztuk": q, "Pełny": "tak" if c.full else "nie",
                            "Wymiary kartonu (cm)": f"{c.dims.length_cm:g}×{c.dims.width_cm:g}×{c.dims.height_cm:g}" if c.dims else "",
                            "Masa (kg)": f"{c.weight_kg:.1f}" if c.weight_kg is not None else "",
                            "Szacowane": "tak" if c.estimated else "",
                        })
            if not pal.layers:
                continue
        for c in g.pack.unplaceable:
            for sku, q in c.contents.items():
                rows.append({"Magazyn": fc, "Paleta": "", "Warstwa": "", "Karton": "", "SKU": sku, "Nazwa": "",
                             "Sztuk": q, "Pełny": "", "Wymiary kartonu (cm)": "brak wymiarów", "Masa (kg)": "", "Szacowane": ""})
    return rows


def summary_rows(ev: PlanEval) -> list[dict]:
    rows = []
    for fc, g in ev.groups.items():
        for pi, pal in enumerate(g.pallets, 1):
            rows.append({
                "Magazyn": fc, "Paleta": pi, "Kartony": len(pal.cartons), "Sztuk": pal.units,
                "Wysokość ładunku (cm)": f"{pal.load_height_cm:.0f}", "Masa brutto (kg)": f"{pal.gross_weight_kg:.0f}",
                "Wypełnienie objętości": f"{pal.fill_ratio:.0%}", "Wypełnienie wysokości": f"{pal.height_ratio:.0%}",
                "Uwagi": "; ".join(pal.issues),
            })
    return rows


def to_csv(rows: list[dict], delimiter: str = ";") -> str:
    if not rows:
        return ""
    buf = io.StringIO()
    w = csv.DictWriter(buf, fieldnames=list(rows[0].keys()), delimiter=delimiter)
    w.writeheader()
    w.writerows(rows)
    return buf.getvalue()


def write_xlsx(path: Path | str, sheets: dict[str, list[dict]]) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    first = True
    for name, rows in sheets.items():
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = name[:31]
        if not rows:
            ws.append(["(brak danych)"])
            continue
        cols = list(rows[0].keys())
        ws.append(cols)
        for r in rows:
            ws.append([r.get(c, "") for c in cols])
        for i, c in enumerate(cols, 1):
            width = max(len(str(c)), *(len(str(r.get(c, ""))) for r in rows)) + 2
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = min(60, width)
    wb.save(str(path))


def export_plan(ev: PlanEval, out_dir: Path | str) -> dict[str, Path]:
    out_dir = Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    pid = ev.plan.plan_id
    files: dict[str, Path] = {}
    p = out_dir / f"{pid}_send_to_amazon.csv"
    p.write_text(to_csv(send_to_amazon_rows(ev), ","), encoding="utf-8")
    files["send_to_amazon"] = p
    x = out_dir / f"{pid}_plan.xlsx"
    write_xlsx(x, {"Pozycje": send_to_amazon_rows(ev), "Pakowanie": packing_rows(ev), "Palety": summary_rows(ev)})
    files["xlsx"] = x
    return files
