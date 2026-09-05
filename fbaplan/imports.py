"""Importers for Seller Central reports and for the catalog workbook.

* Fee preview report ("Podgląd opłat" / "Fee Preview"): packaged unit dimensions,
  weight, size tier and merchant SKU per listing -> ``data/products.csv``.
* FBA inventory reports ("Zarządzaj zapasami FBA" / "Manage FBA Inventory",
  "Uzupełnij zapasy" / "Restock Inventory"): available, inbound, sales ->
  ``data/stock.csv`` (one row per merchant SKU, canonical SKU resolved).
* Catalog workbook (``katalog.xlsx`` exported by the tool, filled in Excel).

Listings are matched to canonical SKUs in this order: merchant SKU already known
in the catalog -> listing title through the alias table (1 354 historical
titles) -> feature key of the title. Newly seen merchant SKUs are remembered on
the product (``msku_werhe`` / ``msku_werkon``).

Report files may be TSV/CSV (UTF-8, UTF-8-BOM, cp1250) or XLSX; headers may be
English machine names (``longest-side``) or localised labels.
"""
from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Optional

from .catalog.store import Resolver, brand_from_name, save_products
from .models import CartonSpec, Dims, Product


# --------------------------------------------------------------------------- #
# Generic table reading
# --------------------------------------------------------------------------- #
def _norm_header(h: str) -> str:
    h = (h or "").strip().lower().replace("_", "-")
    h = re.sub(r"[\s]+", "-", h)
    return h


def read_table(path: Path | str, data: Optional[bytes] = None) -> list[dict[str, str]]:
    """Read TSV/CSV/XLSX into a list of dicts with normalised header keys."""
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data) if data is not None else str(path), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header = [_norm_header(str(c)) if c is not None else "" for c in next(rows, [])]
        out = []
        for r in rows:
            if r is None or all(c in (None, "") for c in r):
                continue
            out.append({header[i]: ("" if c is None else str(c)) for i, c in enumerate(r) if i < len(header) and header[i]})
        return out
    raw = data if data is not None else path.read_bytes()
    text = None
    for enc in ("utf-8-sig", "utf-8", "cp1250", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    assert text is not None
    sample = text[:5000]
    delim = "\t" if sample.count("\t") >= sample.count(";") and sample.count("\t") >= sample.count(",") else (";" if sample.count(";") > sample.count(",") else ",")
    reader = csv.reader(io.StringIO(text), delimiter=delim)
    header = [_norm_header(h) for h in next(reader, [])]
    out = []
    for r in reader:
        if not any(x.strip() for x in r):
            continue
        out.append({header[i]: (r[i].strip() if i < len(r) else "") for i in range(len(header)) if header[i]})
    return out


FIELD_ALIASES = {
    "msku": ["sku", "seller-sku", "merchant-sku", "sku-sprzedawcy", "msku", "händler-sku", "haendler-sku", "sku-sprzedającego"],
    "name": ["product-name", "item-name", "title", "nazwa-produktu", "nazwa", "produktname", "product"],
    "asin": ["asin"],
    "fnsku": ["fnsku"],
    "brand": ["brand", "marka", "marke"],
    "longest": ["longest-side", "najdłuższy-bok", "najdluzszy-bok", "längste-seite", "laengste-seite", "length", "długość", "dlugosc"],
    "median": ["median-side", "średni-bok", "sredni-bok", "mittlere-seite", "width", "szerokość", "szerokosc"],
    "shortest": ["shortest-side", "najkrótszy-bok", "najkrotszy-bok", "kürzeste-seite", "kuerzeste-seite", "height", "wysokość", "wysokosc"],
    "dim_unit": ["unit-of-dimension", "jednostka-wymiaru", "jednostka-wymiarów", "maßeinheit"],
    "weight": ["item-package-weight", "package-weight", "waga-opakowania", "waga", "weight", "gewicht", "verpackungsgewicht"],
    "weight_unit": ["unit-of-weight", "jednostka-wagi", "gewichtseinheit"],
    "size_tier": ["product-size-tier", "klasa-rozmiaru", "size-tier", "größenklasse", "groessenklasse"],
    "available": ["afn-fulfillable-quantity", "available", "dostępne", "dostepne", "fulfillable", "verfügbar", "verfuegbar", "ilość-dostępna"],
    "inbound_shipped": ["afn-inbound-shipped-quantity", "shipped", "wysłane", "wyslane", "inbound"],
    "inbound_working": ["afn-inbound-working-quantity", "working", "w-przygotowaniu"],
    "inbound_receiving": ["afn-inbound-receiving-quantity", "receiving", "przyjmowane", "w-trakcie-przyjmowania"],
    "reserved": ["afn-reserved-quantity", "reserved", "zarezerwowane"],
    "sales_30": ["units-sold-last-30-days", "sales-30d", "sold-30d", "sprzedane-jednostki-w-ciągu-ostatnich-30-dni", "sprzedaż-30-dni", "verkaufte-einheiten-letzte-30-tage", "units-sold-30-days"],
    "sales_90": ["units-sold-last-90-days", "sales-90d", "sold-90d", "sprzedaż-90-dni"],
    "days_of_supply": ["days-of-supply-at-amazon-fulfillment-network", "days-of-supply", "dni-zapasu", "total-days-of-supply-(including-units-from-open-shipments)"],
    "per_unit_volume": ["per-unit-volume", "objętość-jednostkowa"],
}


def _get(row: dict[str, str], key: str) -> str:
    for cand in FIELD_ALIASES[key]:
        v = row.get(cand)
        if v not in (None, ""):
            return v
    return ""


def _num(v: str) -> Optional[float]:
    if v in (None, ""):
        return None
    s = str(v).strip().replace(" ", "").replace(" ", "")
    s = re.sub(r"[^\d,.\-]", "", s)
    if s.count(",") == 1 and s.count(".") == 0:
        s = s.replace(",", ".")
    elif s.count(",") >= 1 and s.count(".") >= 1:
        s = s.replace(",", "")
    try:
        return float(s)
    except ValueError:
        return None


def _to_cm(v: Optional[float], unit: str) -> Optional[float]:
    if v is None:
        return None
    u = (unit or "").lower()
    if u.startswith("in") or u == "inches" or u == "cale":
        return round(v * 2.54, 2)
    if u.startswith("mm") or u.startswith("mil"):
        return round(v / 10.0, 2)
    if u.startswith("m") and not u.startswith("met") and len(u) <= 2:
        return round(v * 100, 2)
    return round(v, 2)


def _to_kg(v: Optional[float], unit: str) -> Optional[float]:
    if v is None:
        return None
    u = (unit or "").lower()
    if u.startswith("lb") or u.startswith("pound") or u.startswith("funt"):
        return round(v * 0.45359237, 3)
    if u.startswith("g") and not u.startswith("kg"):
        return round(v / 1000.0, 3)
    if u.startswith("oz"):
        return round(v * 0.0283495, 3)
    return round(v, 3)


# --------------------------------------------------------------------------- #
# Matching listings to canonical SKUs
# --------------------------------------------------------------------------- #
@dataclass
class ImportReport:
    rows: int = 0
    matched: int = 0
    updated: int = 0
    unmatched: list[str] = field(default_factory=list)   # "msku | name"
    how: dict[str, int] = field(default_factory=dict)
    notes: list[str] = field(default_factory=list)

    def summary(self) -> str:
        parts = [f"wierszy: {self.rows}", f"rozpoznanych: {self.matched}", f"zaktualizowanych: {self.updated}"]
        if self.unmatched:
            parts.append(f"nierozpoznanych: {len(self.unmatched)}")
        return ", ".join(parts) + (" — " + "; ".join(self.notes) if self.notes else "")


class ListingMatcher:
    def __init__(self, products: dict[str, Product], resolver: Resolver):
        self.products = products
        self.resolver = resolver
        self.by_msku = {m: p.sku for p in products.values() for m in p.mskus.values()}

    def match(self, msku: str, name: str) -> tuple[Optional[str], str, str]:
        """Return (canonical sku, brand, how)."""
        brand = brand_from_name(name) or brand_from_name(msku)
        if msku and msku in self.by_msku:
            return self.by_msku[msku], brand, "msku"
        if msku and msku in self.products:
            return msku, brand, "sku"
        if name:
            sku, b, conf, how = self.resolver.resolve(name)
            if sku and (how == "alias" or conf >= 0.4):
                return sku, b or brand, how
        return None, brand, "none"

    def remember_msku(self, sku: str, msku: str, brand: str) -> bool:
        p = self.products.get(sku)
        if not p or not msku:
            return False
        key = brand or ("WERKON" if "werkon" in msku.lower() else "WERHE")
        if p.mskus.get(key) == msku:
            return False
        if msku in p.mskus.values():
            return False
        if key in p.mskus and p.mskus[key] != msku:
            key = key + "-2"
        p.mskus[key] = msku
        if key.split("-")[0] not in p.brands:
            p.brands.append(key.split("-")[0])
        self.by_msku[msku] = sku
        return True


# --------------------------------------------------------------------------- #
# Fee preview -> dimensions
# --------------------------------------------------------------------------- #
def import_fee_preview(path: Path | str, products: dict[str, Product], resolver: Resolver,
                       products_path: Optional[Path | str] = None, data: Optional[bytes] = None,
                       overwrite: bool = False) -> ImportReport:
    rows = read_table(path, data)
    rep = ImportReport(rows=len(rows))
    if rows and not any(_get(rows[0], "longest") for _ in [0]) and not any(_get(r, "longest") for r in rows[:20]):
        rep.notes.append("nie znaleziono kolumny z najdłuższym bokiem (longest-side); nagłówki: " + ", ".join(list(rows[0].keys())[:25]))
    m = ListingMatcher(products, resolver)
    for r in rows:
        msku, name = _get(r, "msku"), _get(r, "name")
        sku, brand, how = m.match(msku, name)
        rep.how[how] = rep.how.get(how, 0) + 1
        if not sku:
            rep.unmatched.append(f"{msku} | {name[:80]}")
            continue
        rep.matched += 1
        p = products[sku]
        changed = m.remember_msku(sku, msku, brand)
        du, wu = _get(r, "dim_unit"), _get(r, "weight_unit")
        L, W, H = (_to_cm(_num(_get(r, k)), du) for k in ("longest", "median", "shortest"))
        w = _to_kg(_num(_get(r, "weight")), wu)
        if L and W and H and (overwrite or p.unit_dims is None):
            p.unit_dims = Dims(L, W, H)
            changed = True
        if w and (overwrite or p.unit_weight_kg is None):
            p.unit_weight_kg = w
            changed = True
        tier = _get(r, "size_tier")
        if tier and f"klasa Amazon: {tier}" not in p.notes:
            p.notes = (p.notes + " | " if p.notes else "") + f"klasa Amazon: {tier}"
            changed = True
        if changed:
            rep.updated += 1
    if products_path:
        save_products(products_path, products.values())
    return rep


# --------------------------------------------------------------------------- #
# Inventory / restock reports -> stock.csv
# --------------------------------------------------------------------------- #
STOCK_COLUMNS = ["sku", "msku", "brand", "fba_available", "fba_inbound", "sales_30d", "sales_90d", "local_stock", "listing_name", "source"]


def import_inventory_report(path: Path | str, products: dict[str, Product], resolver: Resolver,
                            stock_path: Path | str, products_path: Optional[Path | str] = None,
                            data: Optional[bytes] = None, merge: bool = True) -> ImportReport:
    """Read an FBA inventory or restock report and (re)write data/stock.csv.

    With ``merge=True`` existing rows for other merchant SKUs are kept and rows for the
    same merchant SKU are replaced, so an inventory report (stock) and a restock report
    (sales) can be imported one after the other."""
    rows = read_table(path, data)
    rep = ImportReport(rows=len(rows))
    m = ListingMatcher(products, resolver)
    stock_path = Path(stock_path)
    existing: dict[str, dict] = {}
    if merge and stock_path.exists():
        with open(stock_path, newline="", encoding="utf-8-sig") as fh:
            for r in csv.DictReader(fh):
                key = r.get("msku") or r.get("sku")
                if key:
                    existing[key] = r
    src = Path(path).name
    for r in rows:
        msku, name = _get(r, "msku"), _get(r, "name")
        sku, brand, how = m.match(msku, name)
        rep.how[how] = rep.how.get(how, 0) + 1
        if not sku:
            rep.unmatched.append(f"{msku} | {name[:80]}")
            continue
        rep.matched += 1
        m.remember_msku(sku, msku, brand)
        avail = _num(_get(r, "available"))
        inbound = sum(x for x in (_num(_get(r, k)) for k in ("inbound_shipped", "inbound_working", "inbound_receiving")) if x)
        s30, s90 = _num(_get(r, "sales_30")), _num(_get(r, "sales_90"))
        key = msku or sku
        old = existing.get(key, {})
        row = {
            "sku": sku, "msku": msku, "brand": brand,
            "fba_available": int(avail) if avail is not None else (old.get("fba_available") or ""),
            "fba_inbound": int(inbound) if any(_get(r, k) for k in ("inbound_shipped", "inbound_working", "inbound_receiving")) else (old.get("fba_inbound") or ""),
            "sales_30d": int(s30) if s30 is not None else (old.get("sales_30d") or ""),
            "sales_90d": int(s90) if s90 is not None else (old.get("sales_90d") or ""),
            "local_stock": old.get("local_stock") or "",
            "listing_name": name[:120], "source": src,
        }
        existing[key] = row
        rep.updated += 1
    stock_path.parent.mkdir(parents=True, exist_ok=True)
    with open(stock_path, "w", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=STOCK_COLUMNS)
        w.writeheader()
        for row in existing.values():
            w.writerow({k: row.get(k, "") for k in STOCK_COLUMNS})
    if products_path:
        save_products(products_path, products.values())
    return rep


# --------------------------------------------------------------------------- #
# Catalog workbook (Excel round-trip)
# --------------------------------------------------------------------------- #
CATALOG_SHEET_COLUMNS = [
    ("sku", "SKU"), ("name", "Nazwa"), ("family", "Rodzina"), ("priority", "Szt. 12 mies."), ("fc", "Magazyn (przewid.)"),
    ("msku_werhe", "SKU Amazon WERHE"), ("msku_werkon", "SKU Amazon WERKON"),
    ("unit_length_cm", "Sztuka dł. [cm]"), ("unit_width_cm", "Sztuka szer. [cm]"), ("unit_height_cm", "Sztuka wys. [cm]"), ("unit_weight_kg", "Sztuka waga [kg]"),
    ("units_per_carton", "Szt. w kartonie"), ("carton_length_cm", "Karton dł. [cm]"), ("carton_width_cm", "Karton szer. [cm]"), ("carton_height_cm", "Karton wys. [cm]"), ("carton_weight_kg", "Karton waga [kg]"),
    ("pack_group", "Grupa pakowania"), ("max_qty_per_plan", "Maks. szt. w planie"), ("fc_override", "Magazyn ręcznie"), ("active", "Aktywny (1/0)"), ("notes", "Uwagi"),
]


def export_catalog_xlsx(path: Path | str, products: Iterable[Product], priority: dict[str, int], predicted_fc: dict[str, str],
                        estimated_cartons: Optional[dict[str, int]] = None) -> int:
    """Write the catalog as a workbook the client can fill in Excel (most shipped products first)."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    estimated_cartons = estimated_cartons or {}
    wb = Workbook()
    ws = wb.active
    ws.title = "Katalog"
    ws.append([label for _, label in CATALOG_SHEET_COLUMNS])
    fill_need = PatternFill("solid", fgColor="FFF3CD")
    fill_est = PatternFill("solid", fgColor="E7F0FF")
    for c in ws[1]:
        c.font = Font(bold=True)
    n = 0
    for p in sorted(products, key=lambda p: (-priority.get(p.sku, 0), p.sku)):
        row = p.to_row()
        row["priority"] = priority.get(p.sku, 0)
        row["fc"] = predicted_fc.get(p.sku, "")
        if p.carton and p.carton.estimated:
            row["units_per_carton"] = p.carton.units_per_carton
        vals = []
        for key, _ in CATALOG_SHEET_COLUMNS:
            v = row.get(key, "")
            if key in ("unit_length_cm", "unit_width_cm", "unit_height_cm", "unit_weight_kg", "units_per_carton", "carton_length_cm",
                       "carton_width_cm", "carton_height_cm", "carton_weight_kg", "max_qty_per_plan", "priority") and v not in ("", None):
                try:
                    v = float(v)
                    v = int(v) if v == int(v) else v
                except (TypeError, ValueError):
                    pass
            vals.append(v)
        ws.append(vals)
        n += 1
        r = ws.max_row
        for i, (key, _) in enumerate(CATALOG_SHEET_COLUMNS, 1):
            if key in ("unit_length_cm", "unit_width_cm", "unit_height_cm", "unit_weight_kg") and p.active and vals[i - 1] in ("", None):
                ws.cell(row=r, column=i).fill = fill_need
            if key == "units_per_carton" and p.carton and p.carton.estimated:
                ws.cell(row=r, column=i).fill = fill_est
    widths = {1: 16, 2: 60, 3: 12, 4: 12, 5: 12, 6: 22, 7: 22, 21: 40}
    for i in range(1, len(CATALOG_SHEET_COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(i)].width = widths.get(i, 14)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    legend = wb.create_sheet("Legenda")
    for line in [
        "Wypełnij żółte pola: wymiary sztuki w opakowaniu jednostkowym (cm) i waga (kg) — jak w Seller Central (najdłuższy, średni, najkrótszy bok).",
        "Karton zbiorczy: ile sztuk pakujecie do jednego kartonu + wymiary i waga kartonu. Niebieskie pole = liczba sztuk odgadnięta z historii wysyłek, sprawdź.",
        "Grupa pakowania: produkty z różnych grup nie będą proponowane do wspólnej palety (np. „dlugie”, „male”).",
        "Maks. szt. w planie: jeśli danego produktu nie chcecie wysyłać więcej niż X w jednym planie.",
        "Magazyn ręcznie: WRO5 albo XPO1, tylko gdy narzędzie ma się nie zastanawiać.",
        "Aktywny 0 = produkt wycofany, nie będzie proponowany.",
        "Nie zmieniaj kolumny SKU. Po zapisaniu pliku: w narzędziu → Import → „Katalog XLSX”.",
    ]:
        legend.append([line])
    legend.column_dimensions["A"].width = 140
    wb.save(str(path))
    return n


def import_catalog_xlsx(path: Path | str, products: dict[str, Product], products_path: Optional[Path | str] = None,
                        data: Optional[bytes] = None) -> ImportReport:
    rows = read_table(path, data)
    rep = ImportReport(rows=len(rows))
    label_to_key = {_norm_header(label): key for key, label in CATALOG_SHEET_COLUMNS}
    label_to_key.update({key: key for key, _ in CATALOG_SHEET_COLUMNS})

    def g(r: dict, key: str) -> str:
        for h, k in label_to_key.items():
            if k == key and r.get(h) not in (None, ""):
                return str(r[h]).strip()
        return ""

    for r in rows:
        sku = g(r, "sku")
        if not sku:
            continue
        p = products.get(sku)
        if p is None:
            p = Product(sku=sku, name=g(r, "name") or sku)
            products[sku] = p
            rep.notes.append(f"nowy produkt {sku}")
        rep.matched += 1
        before = p.to_row()
        if g(r, "name"):
            p.name = g(r, "name")
        if g(r, "family"):
            p.family = g(r, "family")
        mskus = {}
        if g(r, "msku_werhe"):
            mskus["WERHE"] = g(r, "msku_werhe")
        if g(r, "msku_werkon"):
            mskus["WERKON"] = g(r, "msku_werkon")
        if mskus:
            p.mskus.update(mskus)
            p.brands = sorted(set(p.brands) | set(mskus))
        d = Dims.parse(g(r, "unit_length_cm"), g(r, "unit_width_cm"), g(r, "unit_height_cm"))
        if d:
            p.unit_dims = d
        w = _num(g(r, "unit_weight_kg"))
        if w:
            p.unit_weight_kg = w
        upc = _num(g(r, "units_per_carton"))
        if upc and upc >= 1:
            cd = Dims.parse(g(r, "carton_length_cm"), g(r, "carton_width_cm"), g(r, "carton_height_cm"))
            cw = _num(g(r, "carton_weight_kg"))
            p.carton = CartonSpec(int(upc), cd, cw, estimated=False)
        if g(r, "pack_group"):
            p.pack_group = g(r, "pack_group")
        mq = _num(g(r, "max_qty_per_plan"))
        if mq:
            p.max_qty_per_plan = int(mq)
        fo = g(r, "fc_override").upper()
        p.fc_override = fo or None
        act = g(r, "active")
        if act != "":
            p.active = act.strip().lower() not in ("0", "nie", "false", "n", "no")
        if g(r, "notes"):
            p.notes = g(r, "notes")
        if p.to_row() != before:
            rep.updated += 1
    if products_path:
        save_products(products_path, products.values())
    return rep
